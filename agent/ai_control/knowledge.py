from __future__ import annotations

import contextlib
import csv
import hashlib
import html
import json
import mimetypes
import os
import re
import unicodedata
from pathlib import Path
from typing import Any

from werkzeug.utils import secure_filename

from agent.ai_control.models import AIKnowledgeSourceModel

ANSI_RE = re.compile(r"\x1b\[[0-?]*[ -/]*[@-~]")
MULTI_BLANK_RE = re.compile(r"\n{3,}")


def knowledge_root() -> Path:
    root = Path(os.environ.get("AGENT_AI_DATA_DIR", "ai_control_data")).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    for name in ("originals", "extracted", "clean", "documents"):
        (root / name).mkdir(parents=True, exist_ok=True)
    return root


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def save_upload(file_storage) -> AIKnowledgeSourceModel:
    original_name = file_storage.filename or "upload"
    safe_name = secure_filename(original_name) or "upload"
    ext = Path(safe_name).suffix.lower()
    root = knowledge_root()
    tmp_path = root / "originals" / f"upload-{os.getpid()}-{safe_name}"
    file_storage.save(tmp_path)
    digest = sha256_file(tmp_path)
    stored_name = f"{digest}{ext}"
    final_path = root / "originals" / stored_name
    if final_path.exists():
        tmp_path.unlink(missing_ok=True)
    else:
        tmp_path.replace(final_path)

    mime_type = file_storage.mimetype or mimetypes.guess_type(original_name)[0] or "application/octet-stream"
    size_bytes = final_path.stat().st_size
    return AIKnowledgeSourceModel.create(
        title=Path(original_name).stem or original_name,
        original_filename=original_name,
        stored_filename=stored_name,
        mime_type=mime_type,
        extension=ext,
        size_bytes=size_bytes,
        sha256=digest,
        raw_path=str(final_path),
        metadata_json=json.dumps({"upload_name": original_name}, ensure_ascii=False),
    )


def _extract_plain_text(path: Path) -> tuple[str, str, dict[str, Any]]:
    return path.read_text(encoding="utf-8", errors="replace"), "text", {}


def _extract_json(path: Path) -> tuple[str, str, dict[str, Any]]:
    data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    return json.dumps(data, ensure_ascii=False, indent=2), "json", {"root_type": type(data).__name__}


def _extract_csv(path: Path) -> tuple[str, str, dict[str, Any]]:
    lines: list[str] = []
    rows = 0
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.reader(handle)
        for row in reader:
            rows += 1
            lines.append(" | ".join(str(value) for value in row))
    return "\n".join(lines), "csv", {"rows": rows}


def _extract_html(path: Path) -> tuple[str, str, dict[str, Any]]:
    text = path.read_text(encoding="utf-8", errors="replace")
    text = re.sub(r"(?is)<script.*?>.*?</script>", " ", text)
    text = re.sub(r"(?is)<style.*?>.*?</style>", " ", text)
    text = re.sub(r"(?s)<[^>]+>", "\n", text)
    return html.unescape(text), "html", {}


def _extract_pdf(path: Path) -> tuple[str, str, dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF extraction requires optional package 'pypdf'") from exc
    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n\n".join(pages), "pypdf", {"pages": len(reader.pages)}


def _extract_docx(path: Path) -> tuple[str, str, dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX extraction requires optional package 'python-docx'") from exc
    document = Document(str(path))
    paragraphs = [p.text for p in document.paragraphs]
    tables = []
    for table in document.tables:
        for row in table.rows:
            tables.append(" | ".join(cell.text for cell in row.cells))
    content = "\n".join(paragraphs + (["", "## Tables", *tables] if tables else []))
    return content, "python-docx", {"paragraphs": len(paragraphs), "table_rows": len(tables)}


def _extract_xlsx(path: Path) -> tuple[str, str, dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX extraction requires optional package 'openpyxl'") from exc
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    output: list[str] = []
    sheet_meta: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        output.append(f"# Sheet: {sheet.title}")
        row_count = 0
        for row in sheet.iter_rows(values_only=True):
            row_count += 1
            output.append(" | ".join("" if value is None else str(value) for value in row))
        sheet_meta.append({"name": sheet.title, "rows": row_count})
        output.append("")
    workbook.close()
    return "\n".join(output), "openpyxl", {"sheets": sheet_meta}


def extract_file(path: Path) -> tuple[str, str, dict[str, Any]]:
    ext = path.suffix.lower()
    if ext in {
        ".txt",
        ".md",
        ".log",
        ".py",
        ".js",
        ".ts",
        ".tsx",
        ".jsx",
        ".sql",
        ".xml",
        ".yaml",
        ".yml",
        ".ini",
        ".cfg",
        ".conf",
    }:
        return _extract_plain_text(path)
    if ext == ".json":
        return _extract_json(path)
    if ext in {".csv", ".tsv"}:
        return _extract_csv(path)
    if ext in {".html", ".htm"}:
        return _extract_html(path)
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext in {".xlsx", ".xlsm"}:
        return _extract_xlsx(path)
    raise RuntimeError(f"No extractor enabled for file type '{ext or 'unknown'}'")


def clean_text(content: str) -> tuple[str, dict[str, Any]]:
    original_chars = len(content)
    content = ANSI_RE.sub("", content)
    content = content.replace("\x00", "")
    content = content.replace("\r\n", "\n").replace("\r", "\n")
    content = unicodedata.normalize("NFC", content)
    lines = [line.rstrip() for line in content.split("\n")]
    content = "\n".join(lines)
    content = MULTI_BLANK_RE.sub("\n\n", content).strip()
    return content, {
        "original_chars": original_chars,
        "clean_chars": len(content),
        "removed_chars": max(0, original_chars - len(content)),
    }


def build_reference_document(
    row: AIKnowledgeSourceModel, clean_content: str, extraction_meta: dict[str, Any]
) -> str:
    provenance = {
        "source_file": row.original_filename,
        "sha256": row.sha256,
        "mime_type": row.mime_type,
        "size_bytes": row.size_bytes,
        "extract_method": row.extract_method,
        "extraction": extraction_meta,
    }
    return (
        f"# {row.title}\n\n"
        "## Provenance\n\n"
        "```json\n"
        f"{json.dumps(provenance, ensure_ascii=False, indent=2)}\n"
        "```\n\n"
        "## Clean reference content\n\n"
        f"{clean_content}\n"
    )


def inspect_and_process(row: AIKnowledgeSourceModel) -> AIKnowledgeSourceModel:
    source_path = Path(row.raw_path)
    text, method, extraction_meta = extract_file(source_path)
    clean, cleaning_meta = clean_text(text)
    root = knowledge_root()
    extracted_path = root / "extracted" / f"{row.sha256}.txt"
    clean_path = root / "clean" / f"{row.sha256}.txt"
    document_path = root / "documents" / f"{row.sha256}.md"
    extracted_path.write_text(text, encoding="utf-8")
    clean_path.write_text(clean, encoding="utf-8")

    row.extract_method = method
    row.extracted_path = str(extracted_path)
    row.clean_path = str(clean_path)
    row.char_count = len(clean)
    row.line_count = clean.count("\n") + (1 if clean else 0)
    row.inspection_status = "Passed"
    row.status = "Processed"
    metadata = {
        "extraction": extraction_meta,
        "cleaning": cleaning_meta,
        "original_exists": source_path.exists(),
    }
    row.metadata_json = json.dumps(metadata, ensure_ascii=False, default=str)
    row.modified_at = __import__("datetime").datetime.now()
    row.save()

    reference_document = build_reference_document(row, clean, extraction_meta)
    document_path.write_text(reference_document, encoding="utf-8")
    row.document_path = str(document_path)
    row.status = "ReadyForReview"
    row.save()
    return row


def mark_failed(row: AIKnowledgeSourceModel, error: Exception) -> AIKnowledgeSourceModel:
    metadata = {}
    with contextlib.suppress(json.JSONDecodeError):
        metadata = json.loads(row.metadata_json or "{}")
    metadata["processing_error"] = str(error)
    row.metadata_json = json.dumps(metadata, ensure_ascii=False)
    row.inspection_status = "Failed"
    row.status = "Failed"
    row.save()
    return row


def preview(row: AIKnowledgeSourceModel, variant: str = "clean") -> str:
    path = {
        "raw": row.raw_path,
        "extracted": row.extracted_path,
        "clean": row.clean_path,
        "document": row.document_path,
    }.get(variant)
    if not path:
        return ""
    p = Path(path)
    if variant == "raw" or p.suffix.lower() not in {".txt", ".md", ".json", ".csv"}:
        return ""
    return p.read_text(encoding="utf-8", errors="replace")
